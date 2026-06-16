#!/usr/bin/env python3
"""Build the Integrador Lite Word manual from its Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "manual-integrador-lite.md"
OUTPUT = ROOT / "manual-integrador-lite.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
CODE_BG = "F3F5F7"
TEXT = RGBColor(31, 41, 55)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def keep_lines(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepLines"))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_end])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, color, before, after in (
        ("Title", 25, BLUE, 0, 12),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, DARK_BLUE, 10, 5),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.25)
        style.paragraph_format.left_indent = Inches(0.25)
        style.paragraph_format.first_line_indent = Inches(-0.14)
        style.paragraph_format.space_after = Pt(2.5)

    if "Code Block" not in [style.name for style in document.styles]:
        code = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = document.styles["Code Block"]
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.font.color.rgb = RGBColor(30, 41, 59)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.1)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    header.text = "INTEGRADOR LITE  |  MANUAL OPERATIVO Y TÉCNICO"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_inline_markdown(paragraph, text: str) -> None:
    token_pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_cover(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72)
    run = p.add_run("INTEGRADOR LITE")
    run.font.name = "Calibri"
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(16)
    run = title.add_run("Manual Integral")
    run.font.name = "Calibri"
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Integración multi-cliente HubSpot → Treble").bold = True

    rule = document.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.space_before = Pt(18)
    rule.paragraph_format.space_after = Pt(18)
    run = rule.add_run("OPERACIÓN  •  CONFIGURACIÓN  •  SOPORTE  •  DESARROLLO")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    meta = document.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.style = "Table Grid"
    data = (
        ("Versión", "Rama lite"),
        ("Fecha", "15 de junio de 2026"),
        ("Audiencia", "Administradores, soporte y desarrolladores"),
        ("Fuente", "Código y pruebas ejecutables de Integrador Lite"),
    )
    for row, pair in zip(meta.rows, data):
        for index, value in enumerate(pair):
            cell = row.cells[index]
            set_cell_margins(cell)
            cell.text = value
            if index == 0:
                shade(cell, LIGHT_BLUE)
                cell.paragraphs[0].runs[0].bold = True
    set_repeat_table_header(meta.rows[0])

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(22)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Documento independiente del manual del integrador completo.")
    run.italic = True
    run.font.size = Pt(9)

    document.add_page_break()


def add_toc(document: Document, markdown: str) -> None:
    heading = document.add_paragraph("Contenido", style="Heading 1")
    keep_with_next(heading)
    intro = document.add_paragraph(
        "Las secciones conservan numeración estable y estilos de título para facilitar "
        "la navegación en Word y Google Docs."
    )
    intro.paragraph_format.space_after = Pt(8)
    for line in markdown.splitlines():
        if not line.startswith("## "):
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.12)
        paragraph.paragraph_format.space_after = Pt(2.5)
        run = paragraph.add_run(line[3:])
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    document.add_page_break()


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = len(rows[0])
    table = document.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        shade(cell, LIGHT_BLUE)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        add_inline_markdown(paragraph, value)
        for run in paragraph.runs:
            run.bold = True
            run.font.size = Pt(9)
    set_repeat_table_header(table.rows[0])

    for values in rows[1:]:
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_markdown(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(8.6)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    keep_lines(paragraph)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(7.8)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CODE_BG)
    p_pr.append(shd)


def add_image(document: Document, alt: str, relative_path: str) -> None:
    image_path = ROOT / relative_path
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(6.72))
    inline._inline.docPr.set("descr", alt)
    inline._inline.docPr.set("title", alt)
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run(alt)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(75, 85, 99)


def parse_markdown(document: Document, text: str) -> None:
    lines = text.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []

    while index < len(lines):
        line = lines[index]

        if line.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        image = re.fullmatch(r"!\[(.+?)\]\((.+?)\)", line.strip())
        if image:
            add_image(document, image.group(1), image.group(2))
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|(?:\s*:?-+:?\s*\|)+$", lines[index + 1]):
            rows: list[list[str]] = []
            rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(document, rows)
            continue

        if line.startswith("# "):
            index += 1
            continue
        if line.startswith("## "):
            paragraph = document.add_paragraph(line[3:], style="Heading 1")
            keep_with_next(paragraph)
        elif line.startswith("### "):
            paragraph = document.add_paragraph(line[4:], style="Heading 2")
            keep_with_next(paragraph)
        elif line.startswith("#### "):
            paragraph = document.add_paragraph(line[5:], style="Heading 3")
            keep_with_next(paragraph)
        elif re.match(r"^\d+\.\s+", line):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.24)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_after = Pt(2.5)
            add_inline_markdown(paragraph, line)
        elif line.startswith("- [ ] "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, "☐ " + line[6:])
        elif line.startswith("- [x] "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, "☒ " + line[6:])
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, line[2:])
        elif line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.right_indent = Inches(0.15)
            add_inline_markdown(paragraph, line[2:])
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
            p_pr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), LIGHT_BLUE)
            p_pr.append(shd)
        elif line.strip() == "---":
            document.add_paragraph("────────────────────────────────────────────────────────")
        elif line.strip():
            paragraph = document.add_paragraph()
            add_inline_markdown(paragraph, line)

        index += 1


def set_core_properties(document: Document) -> None:
    props = document.core_properties
    props.title = "Manual Integral de Integrador Lite"
    props.subject = "Operación y arquitectura HubSpot a Treble"
    props.author = "Equipo Integrador"
    props.keywords = "Integrador Lite, HubSpot, Treble, webhooks, reglas, records"
    props.comments = "Generado reproduciblemente desde docs/manual-integrador-lite.md"


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    document = Document()
    configure_document(document)
    set_core_properties(document)
    add_cover(document)
    add_toc(document, markdown)
    parse_markdown(document, markdown)

    final_section = document.add_section(WD_SECTION.CONTINUOUS)
    final_section.header.is_linked_to_previous = True
    final_section.footer.is_linked_to_previous = True

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
